# ============================================================
# Quality audit — time & location violations (GitHub Actions)
# Uses submission metadata only: KSA local time, duration, GPS, distance.
# Output: xlsx (Visits + Summary) emailed with Arabic summary tables.
# Secrets: ZENPUT_API_TOKEN, GMAIL_APP_PASSWORD
# Env: MODE = auto (send) | test (no email)
#      REPORT_DATE = "" (yesterday KSA) or YYYY-MM-DD
#      DAYS = number of days ending on REPORT_DATE (default 1)
# ============================================================
import os, sys, smtplib, requests, json, math
from collections import Counter
from datetime import datetime, timedelta
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email.utils import formataddr
from email import encoders
import pytz
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

API_TOKEN    = os.environ.get("ZENPUT_API_TOKEN", "")
GMAIL_USER   = "aof.group.auto@gmail.com"
GMAIL_PASS   = os.environ.get("GMAIL_APP_PASSWORD", "")
SENDER_NAME  = "Business Intelligence"
TO_EMAILS    = ["o.salahaddin@aofgroup.com"]
CC_EMAILS    = ["a.alsalem@aofgroup.com"]
MODE         = os.environ.get("MODE", "auto").strip().lower()
REPORT_DATE  = os.environ.get("REPORT_DATE", "").strip()
TEMPLATE_IDS = [820340, 821186, 821187, 835723, 709112,
                878309, 1031151, 1059192, 1105384, 1085784]
DAYS         = int(os.environ.get("DAYS", "1") or 1)
ROLE_FILTER  = "Quality Inspectors"   # None = all roles
BASE         = "https://www.zenput.com/api/v3/submissions/"
HEADERS      = {"X-API-TOKEN": API_TOKEN}
TZ           = pytz.timezone("Asia/Riyadh")
PAGE         = 100
LINK_BASE    = "https://ops-execution.crunchtime.com/reports/#form_id/"

# ---- violation thresholds (proposed — adjust) ----
MAX_DISTANCE    = 0.5     # in Zenput's distance unit (see printout) — farther = "Far from branch"
MAX_DURATION_H  = 4       # form open longer than this = "Too long"
MIN_DURATION_M  = 10      # shorter than this = "Too short"
SAME_SPOT_DEC   = 4       # lat/lng rounded to 4 decimals (~11 m) for "same spot, different branches"

# ---------------- helpers ----------------
def parse_local(s):
    """date_submitted_local: naive KSA wall time."""
    if not s: return None
    try:
        d = datetime.fromisoformat(str(s).replace("Z", ""))
        return d.replace(tzinfo=None)
    except Exception: return None

def parse_utc(s):
    """date_submitted / date_started: naive values are UTC."""
    if not s: return None
    try:
        d = datetime.fromisoformat(str(s).replace("Z", "+00:00"))
        if d.tzinfo is None: d = pytz.utc.localize(d)
        return d.astimezone(TZ).replace(tzinfo=None)
    except Exception: return None

def hms(seconds):
    if seconds is None or (isinstance(seconds, float) and math.isnan(seconds)): return ""
    sign = "-" if seconds < 0 else ""
    s = int(abs(round(seconds)))
    return f"{sign}{s // 3600:02d}:{(s % 3600) // 60:02d}:{s % 60:02d}"

def first(d, *keys):
    for k in keys:
        if d.get(k) not in (None, ""): return d.get(k)
    return None

def num(v):
    try: return float(v)
    except Exception: return None

def flatten(d, prefix=""):
    out = {}
    if isinstance(d, dict):
        for k, v in d.items(): out.update(flatten(v, f"{prefix}{k}."))
    else:
        out[prefix[:-1]] = d
    return out

def fetch(tid, since_utc):
    out, start = [], 0
    while True:
        r = requests.get(BASE, headers=HEADERS, params={
            "form_template_id": tid, "limit": PAGE, "start": start}, timeout=60)
        if r.status_code != 200:
            print(f"[{tid}] HTTP {r.status_code}: {r.text[:200]}"); break
        data = r.json().get("data", [])
        if not data: break
        out.extend(data)
        dts = [d for d in (parse_utc((s.get("smetadata") or {}).get("date_submitted")) for s in data) if d]
        if len(data) < PAGE or (dts and min(dts) < since_utc): break
        start += PAGE
    return out

# ---------------- fetch ----------------
if not API_TOKEN:
    sys.exit("ZENPUT_API_TOKEN is not set")
now = datetime.now(TZ).replace(tzinfo=None)
end_day = datetime.strptime(REPORT_DATE, "%Y-%m-%d").date() if REPORT_DATE else (now - timedelta(days=1)).date()
start_day = end_day - timedelta(days=DAYS - 1)
since = datetime.combine(start_day, datetime.min.time()) - timedelta(hours=12)   # fetch margin
until = datetime.combine(end_day, datetime.min.time()) + timedelta(days=1)
period = f"{start_day}" if DAYS == 1 else f"{start_day} - {end_day}"
print(f"Report period: {period} (KSA)")
subs_all = []
for tid in TEMPLATE_IDS:
    subs = fetch(tid, since)
    subs = [s for s in subs if (parse_utc((s.get("smetadata") or {}).get("date_submitted")) or since) >= since]
    for s in subs: s["_tid"] = tid
    subs_all.extend(subs)
    print(f"[{tid}] {len(subs)} submissions fetched")

# ---------------- rows ----------------
rows = []
for s in subs_all:
    sm = s.get("smetadata") or {}
    role = (sm.get("user_role") or {}).get("name")
    if ROLE_FILTER and role != ROLE_FILTER: continue
    loc = sm.get("location") or {}
    submitted = parse_local(sm.get("date_submitted_local")) or parse_utc(sm.get("date_submitted"))
    ttc_ms = num(sm.get("time_to_complete"))
    ttc_s = ttc_ms / 1000 if ttc_ms is not None else None
    started = submitted - timedelta(seconds=ttc_s) if (submitted and ttc_s is not None) else None
    lat = num(first(sm, "lat", "latitude"))
    lng = num(first(sm, "lng", "lon", "long", "longitude"))
    dist = num(sm.get("distance_to_account"))
    gps_on = lat is not None or dist is not None
    rows.append({
        "Template": s["_tid"],
        "Inspector": (sm.get("created_by") or {}).get("display_name"),
        "Branch": " ".join(str(loc.get("name") or "").split()),
        "GPS": "On" if gps_on else "OFF",
        "Lat": lat, "Lng": lng,
        "Distance to branch": round(dist, 2) if dist is not None else None,
        "_start": started, "_sub": submitted, "_ttc": ttc_s,
        "Link": LINK_BASE + str(s.get("legacy_submission_id") or ""),
    })

df = pd.DataFrame(rows)
if not df.empty:
    df = df[(df["_sub"] >= datetime.combine(start_day, datetime.min.time())) & (df["_sub"] < until)]
if df.empty:
    print("No quality visits in the period - nothing to send."); sys.exit(0)
df = df.sort_values(["Inspector", "_sub"]).reset_index(drop=True)

# previous submission by same inspector
df["_prev_sub"]    = df.groupby("Inspector")["_sub"].shift(1)
df["Prev Branch"]  = df.groupby("Inspector")["Branch"].shift(1)
df["_gap"]         = (df["_start"] - df["_prev_sub"]).dt.total_seconds()

# same spot used for different branches (same inspector)
df["_spot"] = df.apply(lambda r: f"{round(r['Lat'], SAME_SPOT_DEC)},{round(r['Lng'], SAME_SPOT_DEC) if r['Lng'] is not None and not pd.isna(r['Lng']) else ''}"
                       if r["Lat"] is not None and not pd.isna(r["Lat"]) else None, axis=1)
spot_branches = df.dropna(subset=["_spot"]).groupby(["Inspector", "_spot"])["Branch"].nunique()
multi_spot = set(spot_branches[spot_branches > 1].index)

def flags(r):
    f = []
    if r["GPS"] == "OFF": f.append("GPS off")
    if r["Distance to branch"] is not None and not pd.isna(r["Distance to branch"]) and r["Distance to branch"] > MAX_DISTANCE:
        f.append("Far from branch")
    if r["_spot"] and (r["Inspector"], r["_spot"]) in multi_spot:
        f.append("Same spot, different branches")
    if r["_ttc"] is not None and r["_ttc"] > MAX_DURATION_H * 3600: f.append("Too long")
    if r["_ttc"] is not None and r["_ttc"] < MIN_DURATION_M * 60:  f.append("Too short")
    if not pd.isna(r["_gap"]) and r["_gap"] < 0 and r["Prev Branch"] != r["Branch"]:
        f.append("Overlaps previous branch")
    return ", ".join(f)

df["Violations"] = df.apply(flags, axis=1)

out = pd.DataFrame({
    "Inspector": df["Inspector"],
    "Branch": df["Branch"],
    "Date": df["_sub"].dt.strftime("%Y-%m-%d"),
    "Duration (HH:MM:SS)": df["_ttc"].apply(hms),
    "GPS": df["GPS"],
    "Violations": df["Violations"],
})

# ---------------- summary ----------------
print("\n" + "=" * 90 + "\nVIOLATIONS BY INSPECTOR\n" + "=" * 90)
kinds = ["GPS off", "Far from branch", "Same spot, different branches", "Too long", "Too short", "Overlaps previous branch"]
summ = pd.DataFrame({k: out.groupby("Inspector")["Violations"].apply(lambda s: s.str.contains(k, regex=False).sum()) for k in kinds})
summ.insert(0, "Visits", out.groupby("Inspector").size())
summ["Visits with any violation"] = out.groupby("Inspector")["Violations"].apply(lambda s: (s != "").sum())
print(summ.to_string())

# ---------------- xlsx ----------------
fname = f"quality_time_location_audit_{end_day}.xlsx"
with pd.ExcelWriter(fname, engine="openpyxl") as w:
    out.to_excel(w, index=False, sheet_name="Visits")
    summ.reset_index().to_excel(w, index=False, sheet_name="Summary")

wb = load_workbook(fname)
HDR, RED = PatternFill("solid", fgColor="1F3864"), PatternFill("solid", fgColor="F8CBAD")
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
for ws in wb.worksheets:
    for c in ws[1]:
        c.fill, c.font = HDR, Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, min(40, max(len(str(c.value or "")) for c in col) + 2))
    for row in ws.iter_rows(min_row=2):
        for c in row: c.border, c.alignment = THIN, Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
ws = wb["Visits"]
cols = {c.value: c.column - 1 for c in ws[1]}
for row in ws.iter_rows(min_row=2):
    if row[cols["Violations"]].value:
        row[cols["Violations"]].fill = RED
    if row[cols["GPS"]].value == "OFF":
        row[cols["GPS"]].fill = RED
wb.save(fname)
print(f"\nSaved {fname}")

# ---------------- email ----------------
KIND_AR = {
    "GPS off": "الموقع مغلق",
    "Far from branch": "بعيد عن الفرع",
    "Same spot, different branches": "نفس الموقع لفروع مختلفة",
    "Too long": "مدة طويلة",
    "Too short": "مدة قصيرة",
    "Overlaps previous branch": "تداخل مع زيارة فرع آخر",
}
TD = "padding:6px 10px;border:1px solid #cccccc;text-align:center"
TH = "padding:8px 10px;background:#1F3864;color:#ffffff;text-align:center"

def cell(v, red=False):
    style = TD + (";color:#C00000;font-weight:bold" if red and v else "")
    return f"<td style='{style}'>{v}</td>"

def build_html():
    total = len(out)
    flagged = int((out["Violations"] != "").sum())
    # table 1: totals by violation type
    t1 = "".join(f"<tr><td style='{TD}'>{KIND_AR[k]}</td>{cell(int(summ[k].sum()), True)}</tr>" for k in kinds)
    # table 2: per inspector
    hdr = "".join(f"<th style='{TH}'>{KIND_AR[k]}</th>" for k in kinds)
    t2 = ""
    for insp, r in summ.sort_values("Visits with any violation", ascending=False).iterrows():
        t2 += (f"<tr><td style='{TD}' dir='ltr'>{insp}</td>{cell(int(r['Visits']))}"
               + "".join(cell(int(r[k]), True) for k in kinds)
               + cell(int(r["Visits with any violation"]), True) + "</tr>")
    return f"""<html><body dir="rtl" style="font-family:Tahoma,Arial,sans-serif;font-size:14px;color:#222;text-align:right">
<p>السلام عليكم ورحمة الله وبركاته،</p>
<p>مرفق لكم تقرير تدقيق الوقت والموقع لزيارات مفتشي الجودة للفترة <span dir="ltr">{period}</span>.</p>
<p>إجمالي الزيارات: <b>{total}</b> &nbsp;|&nbsp; زيارات بها مخالفة: <b style="color:#C00000">{flagged}</b></p>
<h3 style="color:#1F3864;margin:18px 0 8px">ملخص المخالفات حسب النوع</h3>
<table style="border-collapse:collapse">
<tr><th style="{TH}">نوع المخالفة</th><th style="{TH}">العدد</th></tr>{t1}
</table>
<h3 style="color:#1F3864;margin:18px 0 8px">ملخص المخالفات حسب المفتش</h3>
<table style="border-collapse:collapse">
<tr><th style="{TH}">المفتش</th><th style="{TH}">الزيارات</th>{hdr}<th style="{TH}">زيارات بها مخالفة</th></tr>{t2}
</table>
<p style="margin-top:14px;color:#555">معايير المخالفة: الموقع مغلق = لا توجد بيانات موقع؛ بعيد عن الفرع = المسافة أكبر من {MAX_DISTANCE}؛
مدة طويلة = أكثر من {MAX_DURATION_H} ساعات؛ مدة قصيرة = أقل من {MIN_DURATION_M} دقائق؛
تداخل = بدء زيارة قبل تسليم زيارة فرع آخر.</p>
<p>تفاصيل جميع الزيارات في الملف المرفق.</p>
<p>مع خالص التحية،<br>إدارة ذكاء الأعمال</p>
</body></html>"""

def send_email(path):
    if not GMAIL_PASS:
        sys.exit("GMAIL_APP_PASSWORD is not set")
    msg = MIMEMultipart()
    msg["From"] = formataddr((SENDER_NAME, GMAIL_USER))
    msg["To"] = ", ".join(TO_EMAILS)
    msg["Cc"] = ", ".join(CC_EMAILS)
    msg["Subject"] = f"تقرير تدقيق الوقت والموقع - مفتشي الجودة - {period}"
    msg.attach(MIMEText(build_html(), "html", "utf-8"))
    with open(path, "rb") as f:
        part = MIMEBase("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        part.set_payload(f.read())
    encoders.encode_base64(part)
    part.add_header("Content-Disposition", f'attachment; filename="{os.path.basename(path)}"')
    msg.attach(part)
    with smtplib.SMTP_SSL("smtp.gmail.com", 465) as s:
        s.login(GMAIL_USER, GMAIL_PASS)
        s.sendmail(GMAIL_USER, TO_EMAILS + CC_EMAILS, msg.as_string())
    print(f"Email sent to {TO_EMAILS} cc {CC_EMAILS}")

if MODE == "test":
    print("MODE=test -> email skipped")
else:
    send_email(fname)
